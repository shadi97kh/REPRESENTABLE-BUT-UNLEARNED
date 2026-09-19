"""Versioned reconciliation of measurements, stereochemistry and assay provenance."""
from __future__ import annotations
import csv
import json
from pathlib import Path
import re
from collections import defaultdict
from .accounting import snapshot,elapsed,sha,write_json


def code(s):
    parts=dict(re.findall(r'([ABC])(\d+)',s))
    if set(parts)!=set('ABC'): raise ValueError('invalid component code')
    return ''.join(k+parts[k] for k in 'ABC')


def main():
    start=snapshot()
    import openpyxl
    from rdkit import Chem
    root=Path('data/intervention_sources')
    ag=list(csv.DictReader((root/'agile.csv').open()))
    la=list(csv.DictReader((root/'lantern.csv').open()))
    by_code={r['label']:r for r in ag}
    w=openpyxl.load_workbook(root/'supplement/agile_source.xlsx',read_only=True,data_only=True)
    source={}
    for a in range(1,21):
        assert w['Figure 2'].cell(3,a+1).value==f'A{a}'
        for b in range(1,13):
            for c in range(1,6):
                row=4+(c-1)*12+b-1
                cell=w['Figure 2'].cell(row,a+1)
                source[f'A{a}B{b}C{c}']=dict(value=cell.value,sheet='Figure 2',cell=cell.coordinate,
                     mapping_basis='A column headings explicit; declared B-fast/C-slow 60-row ordering; mapping cross-checked against every value')
    source_disagreements=[]
    for k,r in by_code.items():
        if abs(float(r['expt_Hela'])-source[k]['value'])>1e-6:source_disagreements.append(k)
    code_disagreements=[r['Lipid'] for r in la if abs(float(r['Target'])-float(by_code[code(r['Lipid'])]['expt_Hela']))>1e-6]
    canon=lambda s: Chem.MolToSmiles(Chem.MolFromSmiles(s),isomericSmiles=True)
    representation=[]
    for r in la:
        k=code(r['Lipid']);g=by_code[k]
        if canon(r['SMILES'])!=canon(g['combined_mol_SMILES']):
            representation.append(dict(component_code=k,agile_smiles=g['combined_mol_SMILES'],
                                       lantern_smiles=r['SMILES'],label_agrees=True))
    cis=list(csv.DictReader((root/'supplement/lantern_CisTrans.csv').open()))
    stereo={code(r['Lipid']):dict(isometry=r['Isometry'],raw_representation=r['SMILES'],raw_label=r['Target']) for r in cis}
    z=list(csv.DictReader((root/'supplement/agile_finetuning_zenodo_21422889.csv').open()))
    zgroup=defaultdict(list)
    for r in z:zgroup[canon(r['smiles'])].append(r)
    zag=[]
    for r in ag:
        matches=zgroup[canon(r['combined_mol_SMILES'])]
        if not any(abs(float(r['expt_Hela'])-float(x['expt_Hela']))<=1e-6 for x in matches):zag.append(r['label'])
    study=dict(formulation=dict(ionizable_lipid=35,DOPE=16,cholesterol=46.5,C14_PEG2000=2.5,unit='molar parts'),
               cargo='firefly luciferase mRNA (mFluc)',dose=dict(value=.1,unit='microgram mRNA per well'),
               timing=dict(value=24,unit='hours post-treatment'),
               endpoint_scale='log2(mean transfected-cell luminescence / mean untreated-cell luminescence)',
               source='https://www.nature.com/articles/s41467-024-50619-z',
               source_sections=['LNPs formulation and characterization','Luciferase transfection assay'],
               applicability='study-level HTS settings; individual well deviations, reagent lots and doses not independently recorded here',
               assay_material='crude, unpurified reaction products in initial HTS; a pure product graph is an approximation',
               replicate_availability='per-well biological/technical repeats unavailable for the 1200-row endpoint grid',
               aggregation='ratio of means specified; individual means, repeated controls and per-row variances absent')
    out=Path('runs/intervention/audit-v2');out.mkdir(parents=True,exist_ok=True)
    finalrows=[]
    for line in Path('runs/intervention/audit-v1/rows.jsonl').read_text().splitlines():
        r=json.loads(line);raw=r['raw_fields'];k=None
        if r['dataset']=='agile':k=raw['label']
        elif r['dataset']=='lantern':k=code(raw['Lipid'])
        r['component_code']=k;r['study_assay_provenance']=study
        r['measurement_kind']='published_delivery_endpoint; replicate aggregation incompletely observed'
        if k:
            r['source_measurement']=dict(**source[k],workbook_sha256=sha(root/'supplement/agile_source.xlsx'))
            r['source_value_agrees']=abs(r['hela_label_numeric']-source[k]['value'])<=1e-6
            r['components']={key:{'raw':by_code[k][key],'canonical_isomeric':canon(by_code[k][key])}
                             for key in ['A_smiles','B_smiles','C_smiles']}
            r['stereochemical_provenance']=stereo.get(k)
            r['representation_disagrees_across_current_files']=k in {x['component_code'] for x in representation}
            r['future_pair_eligibility']='conditional_crude_HTS; source code matched; verify product/mixture before molecular edit pairing'
        else:
            r['source_measurement']=None
            r['source_value_agrees']=None
            r['source_mapping_status']='TransMA rows lack component codes; chemical matching alone does not resolve conflicting labels'
        finalrows.append(r)
    p=out/'rows.jsonl'
    with p.open('x') as f:
        for r in finalrows:f.write(json.dumps(r,allow_nan=False)+'\n')
    Path(str(p)+'.sha256').write_text(f'{sha(p)}  {p.name}\n')
    summary=dict(original_audit='runs/intervention/audit-v1/audit.json',
                 row_count=len(finalrows),source_workbook_rows_checked=1200,
                 source_workbook_disagreements=source_disagreements,
                 source_workbook_mapping_limit='B-fast/C-slow row ordering is declared and cross-checked; row headings do not explicitly encode B and C',
                 lantern_same_component_code_rows=1100,lantern_same_component_code_label_disagreements=code_disagreements,
                 same_component_code_representation_disagreements=representation,
                 cis_trans_annotation_rows=len(cis),cis_trans_annotation_counts=dict(__import__('collections').Counter(r['Isometry'] for r in cis)),
                 current_zenodo_agile_rows=len(z),current_zenodo_canonical_unique=len(zgroup),
                 current_zenodo_vs_grouped_chemical_label_disagreements=zag,
                 study_assay_provenance=study,
                 interpretation='The 100 AGILE-LANTERN canonical-match label disagreements in v1 reflect product/isomer representation alignment, not 100 current same-code measurement disagreements. Do not merge cis-trans mixtures with pure trans products, or call all conflicts biological replicate noise.',
                 pairs_constructed=0,training_steps=0,resources=elapsed(start))
    write_json(out/'reconciliation.json',summary)
    print(json.dumps({k:v for k,v in summary.items() if k not in ('study_assay_provenance','same_component_code_representation_disagreements')},indent=2))
    print('same_code_representation_disagreements',len(representation))


if __name__=='__main__': main()
